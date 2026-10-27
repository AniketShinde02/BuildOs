import pandas as pd
from openpyxl import load_workbook
import os

file_path = "d:\\Linkedin Automation\\PromptForge_Final_Assets_v4.xlsx"

# The user wants highly structured prompts focusing on face, body structure, and figure
# because the dress will be swapped later. The base outfit should be simple to allow easy swapping.

final_models_data = [
    {
        "Model ID": "BASE_SWAP_01",
        "Category": "Elegant Traditional (Wheatish)",
        "Prompt": "Create a photorealistic image of a beautiful young Indian female fashion model for a premium ethnic wear brand. Age 20–23, Warm wheatish skin, Long silky black hair with middle partition, Soft natural smile, Large brown eyes with minimal kajal. FOCUS ON BODY STRUCTURE: Slim, healthy physique, elegant posture, standing tall with natural proportions. Wearing a simple, form-fitting plain white long Kurti and tight white leggings (designed for easy garment swapping). Luxury Indian fashion catalogue style, beautiful expressive eyes, healthy natural skin with visible realistic skin texture and subtle pores. Highly photorealistic, ultra-detailed, 8K, cinematic soft morning sunlight."
    },
    {
        "Model ID": "BASE_SWAP_02",
        "Category": "Rich Dusky Beauty (Curly Hair)",
        "Prompt": "Create a photorealistic image of a beautiful young Indian female fashion model for a premium ethnic wear brand. Age 22–25, Rich dusky golden skin, Long naturally curly black hair, Strong expressive eyes, Elegant smile. FOCUS ON BODY STRUCTURE: Athletic yet feminine figure, perfect posture, graceful natural stance. Wearing a simple, form-fitting plain white long Kurti and tight white leggings (designed for easy garment swapping). Authentic Indian beauty, healthy glowing skin, sharp but soft facial features. Professional editorial photography, premium boutique catalogue, ultra-detailed, 8K, natural studio lighting."
    },
    {
        "Model ID": "BASE_SWAP_03",
        "Category": "Modern Premium (Layered Hair)",
        "Prompt": "Create a photorealistic image of a beautiful young Indian female fashion model for a premium ethnic wear brand. Age 21–24, Honey wheatish skin, Long layered hair with soft waves, Big brown eyes, Natural makeup. FOCUS ON BODY STRUCTURE: Tall, slender fashion model figure, confident runway posture, standing straight. Wearing a simple, form-fitting plain white long Kurti and tight white leggings (designed for easy garment swapping). Luxury ethnic wear advertisement, confident yet approachable appearance, balanced facial symmetry, no heavy glamour. Commercial fashion photography, 8K, photorealistic, golden hour sunlight."
    },
    {
        "Model ID": "BASE_SWAP_04",
        "Category": "Soft College Girl (Fresh Face)",
        "Prompt": "Create a photorealistic image of a beautiful young Indian female fashion model for a premium ethnic wear brand. Age 19–22, Fresh youthful face, Natural wheatish complexion, Long straight black hair, Cute smile, Bright expressive eyes. FOCUS ON BODY STRUCTURE: Petite, healthy natural physique, relaxed candid posture. Wearing a simple, form-fitting plain white sleeveless Kurti and tight white leggings (designed for easy garment swapping). Lifestyle fashion photography, minimal elegant makeup, well-groomed brows, thick healthy hair. Fresh morning light, highly photorealistic, 8K, natural beauty."
    },
    {
        "Model ID": "BASE_SWAP_05",
        "Category": "Premium Bridal Base (Caramel Skin)",
        "Prompt": "Create a photorealistic image of a beautiful young Indian female fashion model for a premium ethnic wear brand. Age 23–27, Golden caramel skin, Long thick silky hair tied in a loose low bun, Elegant facial features, Natural lips. FOCUS ON BODY STRUCTURE: Hourglass natural proportions, regal and graceful posture, standing elegantly. Wearing a simple, form-fitting plain white long Kurti and tight white leggings (designed for easy garment swapping). High-end Indian designer campaign, museum-quality realism, visible realistic skin texture, soft blush. Editorial luxury fashion photography, ultra-detailed, 8K, cinematic lighting."
    },
    {
        "Model ID": "BASE_SWAP_06",
        "Category": "High-Fashion Editorial (Sharp Features)",
        "Prompt": "Create a photorealistic image of a beautiful young Indian female fashion model for a premium ethnic wear brand. Age 22–26, Warm fair skin, Sharp jawline, High cheekbones, Sleek straight hair tucked behind ears, Fierce confident gaze. FOCUS ON BODY STRUCTURE: Tall, editorial runway figure, striking angular posture. Wearing a simple, form-fitting plain white long Kurti and tight white leggings (designed for easy garment swapping). Premium textile advertisement, authentic Indian beauty, no unrealistic beauty, perfectly balanced facial symmetry. 8K, extremely realistic, flat lay studio lighting."
    },
    {
        "Model ID": "BASE_SWAP_07",
        "Category": "Earthy & Organic Beauty",
        "Prompt": "Create a photorealistic image of a beautiful young Indian female fashion model for a premium ethnic wear brand. Age 24–28, Deep warm olive skin tone, Messy natural wavy hair, Soft neutral expression, Piercing dark eyes. FOCUS ON BODY STRUCTURE: Healthy, curvy natural figure, relaxed and grounded posture. Wearing a simple, form-fitting plain white long Kurti and tight white leggings (designed for easy garment swapping). Pinterest luxury fashion board style, no heavy glamour, highly realistic pores and skin texture. 8K, natural soft window lighting, ultra-detailed."
    },
    {
        "Model ID": "BASE_SWAP_08",
        "Category": "Classic Catalogue (Universal)",
        "Prompt": "Create a photorealistic image of a beautiful young Indian female fashion model for a premium ethnic wear brand. Age 21–25, Medium wheatish skin, Wavy shoulder-length hair, Gentle approachable smile, Subtle kohl. FOCUS ON BODY STRUCTURE: Standard e-commerce mannequin figure, perfectly straight posture facing the camera. Wearing a simple, form-fitting plain white long Kurti and tight white leggings (designed for easy garment swapping). Myntra Luxe catalogue style, clean commercial aesthetic, healthy natural skin. 8K, bright even studio lighting, photorealistic."
    }
]

df_models = pd.DataFrame(final_models_data)

try:
    if os.path.exists(file_path):
        book = load_workbook(file_path)
        
        # Remove old MASTER_MODELS
        if "MASTER_MODELS" in book.sheetnames:
            del book["MASTER_MODELS"]
            
        book.save(file_path)
        
        with pd.ExcelWriter(file_path, engine="openpyxl", mode="a") as writer:
            df_models.to_excel(writer, sheet_name="MASTER_MODELS", index=False)
            
        print("Successfully updated MASTER_MODELS in V4 for easy dress swapping.")
    else:
        print(f"File not found: {file_path}")
except PermissionError:
    print("PERMISSION_ERROR")
except Exception as e:
    print(f"Error: {e}")
