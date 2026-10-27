import pandas as pd
from openpyxl import load_workbook
import os

file_path = "d:\\Linkedin Automation\\PromptForge_Final_Assets_v4.xlsx"

new_models_data = [
    {
        "Model_ID": "NAT001", 
        "Category": "Natural & Radiant (Wavy Hair)", 
        "Prompt": "Cinematic portrait of a gorgeously natural Indian fashion model, radiant glowing complexion, striking confident eyes with subtle kohl, thick voluminous wavy dark hair falling over shoulders. She has a highly photorealistic, expressive face with a gentle confident smile. She is wearing a simple plain white long Kurti with full sleeves and matching trousers. Full body shot, natural relaxed posture. 85mm lens, f/2.8, soft diffused studio lighting, 8k, ultra-detailed high-end fashion editorial."
    },
    {
        "Model_ID": "NAT002", 
        "Category": "Elegant & Confident (Layered Hair)", 
        "Prompt": "Breathtaking photorealistic fashion portrait of a stunning Indian model, elegant and confident posture, natural flawless skin with a subtle golden glow, soft expressive brown eyes, messy but styled layered hair. She is wearing a simple plain white long Kurti with full sleeves. Full body shot, standing with a confident editorial pose. 8k, shot on 35mm, cinematic softbox lighting, highly detailed commercial photography."
    },
    {
        "Model_ID": "NAT003", 
        "Category": "Sharp & Poised (Editorial)", 
        "Prompt": "Ultra-realistic editorial shot of a beautiful Indian fashion model with a sharp jawline, deep expressive eyes, and a confident, poised expression. Natural everyday makeup, luscious dark hair gently blowing in a soft breeze. Wearing a pristine white long Kurti and trousers. Full body shot, high-end commercial fashion aesthetic, soft butterfly lighting, 8k resolution, photorealistic."
    },
    {
        "Model_ID": "NAT004", 
        "Category": "Authentic Beauty (Inviting Gaze)", 
        "Prompt": "Stunningly authentic portrait of a gorgeous Indian female model, soft natural beauty, highly detailed facial features, radiant healthy skin, confident and inviting gaze. Her dark hair is styled in loose natural waves. She wears a simple white long Kurti. Full body shot, relaxed and natural stance. Cinematic lighting, photorealistic, 8k, Vogue editorial style."
    },
    {
        "Model_ID": "NAT005", 
        "Category": "High-Fashion Symmetry (Striking)", 
        "Prompt": "High-fashion photorealistic portrait of an exceptionally beautiful Indian model. She exudes confidence and elegance, featuring striking facial symmetry, warm glowing skin, and captivating eyes with soft eyeliner. Wearing a simple white long Kurti. Full body shot, dynamic yet natural pose, soft even studio lighting, 85mm lens, 8k, commercial quality."
    }
]

df = pd.DataFrame(new_models_data)

try:
    if os.path.exists(file_path):
        # Load the workbook
        book = load_workbook(file_path)
        
        # Remove the old MASTER_MODELS sheet if it exists
        if "MASTER_MODELS" in book.sheetnames:
            del book["MASTER_MODELS"]
        
        # Save the workbook without the old sheet
        book.save(file_path)
        
        # Now append the new sheet
        with pd.ExcelWriter(file_path, engine="openpyxl", mode="a") as writer:
            df.to_excel(writer, sheet_name="MASTER_MODELS", index=False)
            
        print("Successfully updated MASTER_MODELS in V4.")
    else:
        print(f"File not found: {file_path}")
except PermissionError:
    print("PERMISSION_ERROR")
except Exception as e:
    print(f"Error: {e}")
