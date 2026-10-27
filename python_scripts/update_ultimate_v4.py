import pandas as pd
from openpyxl import load_workbook
import os

file_path = "d:\\Linkedin Automation\\PromptForge_Final_Assets_v4.xlsx"

# The 5 specific models provided by the user
ultimate_models_data = [
    {
        "Model Type": "MODEL 1 — Elegant Traditional Girl",
        "Prompt": "Beautiful Indian woman, Age 20–23, Warm wheatish skin, Long silky black hair, Middle partition, Soft smile, Large brown eyes, Minimal kajal, Natural nude lips, Small bindi, Simple earrings, White Chanderi kurta, Matching straight pants, Elegant dupatta, Standing naturally, Luxury boutique photoshoot, Soft morning sunlight, Photorealistic, Commercial catalogue, Premium ethnic fashion"
    },
    {
        "Model Type": "MODEL 2 — Rich Dusky Beauty",
        "Prompt": "Beautiful Indian female model, Age 22–25, Rich dusky golden skin, Long naturally curly black hair, Strong expressive eyes, Minimal makeup, Healthy glowing skin, Sharp but soft facial features, Elegant smile, Ivory Anarkali suit, Traditional jhumkas, Natural standing pose, Luxury Indian clothing campaign, Editorial fashion photography, Premium boutique catalogue, Highly realistic, 8K"
    },
    {
        "Model Type": "MODEL 3 — Modern Premium Boutique Look",
        "Prompt": "Beautiful Indian fashion model, Age 21–24, Honey wheatish skin, Long layered hair, Soft waves, Natural makeup, Big brown eyes, Minimal jewellery, Premium embroidered kurta, Palazzo pants, Graceful walking pose, Outdoor luxury courtyard, Golden hour sunlight, Commercial fashion photography, Pinterest aesthetic, Luxury ethnic wear advertisement, Photorealistic"
    },
    {
        "Model Type": "MODEL 4 — Soft College Girl Look",
        "Prompt": "Beautiful young Indian woman, Age 19–22, Fresh youthful face, Natural wheatish complexion, Long straight black hair, Cute smile, Bright expressive eyes, Minimal makeup, Cotton kurta, Simple earrings, White dupatta, Standing naturally, Lifestyle fashion photography, Luxury ethnic catalogue, Fresh morning light, Natural candid pose, Highly photorealistic"
    },
    {
        "Model Type": "MODEL 5 — Premium Bridal Boutique Style",
        "Prompt": "Beautiful Indian female model, Age 23–27, Golden caramel skin, Long thick silky hair, Elegant facial features, Minimal bridal-inspired makeup, Beautiful eyes, Natural lips, Premium designer embroidered Anarkali, Luxury dupatta, Traditional jewellery, Graceful pose, High-end Indian designer campaign, Editorial luxury fashion photography, Museum-quality realism, Ultra detailed, Premium textile advertisement"
    }
]

# The Master Randomization Prompt and instructions
master_randomizer_data = [
    {
        "Component": "MASTER BASE PROMPT (Use this for dynamic generation)",
        "Content": "Create a photorealistic image of a beautiful young Indian female fashion model for a premium ethnic wear brand. The woman should look naturally attractive rather than unrealistically perfect. She should resemble the kind of models seen in modern Indian fashion campaigns, Pinterest luxury fashion boards, Myntra Luxe, Ajio Luxe, Nykaa Fashion, Aza Fashions and high-end boutique catalogues. Age between 19–27 years. Generate a completely unique Indian face every time."
    },
    {
        "Component": "Mandatory Quality Tags",
        "Content": "beautiful expressive eyes, healthy natural skin, visible realistic skin texture, subtle pores, natural lips, soft blush, minimal elegant makeup, well-groomed brows, thick healthy hair, balanced facial symmetry, confident yet approachable appearance. No heavy glamour. No exaggerated makeup. No unrealistic beauty. The model should have a slim healthy physique with natural proportions. Pose should look effortless like a real professional fashion photoshoot. Luxury Indian fashion catalogue style. Authentic Indian beauty. Professional editorial photography. Highly photorealistic. Ultra detailed. Natural lighting. Luxury fashion campaign. Premium textile advertisement. 8K. Extremely realistic. No AI artifacts."
    },
    {
        "Component": "RANDOMIZE: Hair",
        "Content": "Straight | Soft waves | Loose curls | Ponytail | Low bun | Braid | Half tied"
    },
    {
        "Component": "RANDOMIZE: Expression",
        "Content": "Gentle smile | Confident smile | Neutral | Looking away | Laughing naturally | Soft candid"
    },
    {
        "Component": "RANDOMIZE: Pose",
        "Content": "Standing | Walking | Hand on dupatta | Adjusting earrings | Looking sideways | Holding handbag | Sitting elegantly"
    },
    {
        "Component": "RANDOMIZE: Background",
        "Content": "Luxury boutique | Heritage haveli | Minimal studio | Premium indoor setting | Soft outdoor garden | Luxury hotel corridor | Modern architecture | Neutral textured wall"
    },
    {
        "Component": "UNIVERSAL NEGATIVE PROMPT (ALWAYS USE THIS)",
        "Content": "cartoon, anime, CGI, doll face, plastic skin, fake skin, oversaturated, over sharpened, western facial features, celebrity face, duplicate face, deformed hands, extra fingers, bad anatomy, unrealistic body, blurry, low quality, watermark, logo, text, jpeg artifacts, AI artifacts, ugly face, crossed eyes, exaggerated makeup, unnatural smile, bad teeth, distorted clothing, poor fabric texture"
    },
    {
        "Component": "HOW TO USE THIS (Instruction)",
        "Content": "To get maximum variety and stop getting the exact same AI face: Combine the MASTER BASE PROMPT + Pick 1 Hair + Pick 1 Expression + Pick 1 Pose + Pick 1 Background + Mandatory Quality Tags. Paste the Universal Negative Prompt in the negative section."
    }
]

df_models = pd.DataFrame(ultimate_models_data)
df_randomizer = pd.DataFrame(master_randomizer_data)

try:
    if os.path.exists(file_path):
        book = load_workbook(file_path)
        
        # Remove old MASTER_MODELS if it exists to replace it with the new layout
        if "MASTER_MODELS" in book.sheetnames:
            del book["MASTER_MODELS"]
            
        book.save(file_path)
        
        with pd.ExcelWriter(file_path, engine="openpyxl", mode="a") as writer:
            df_models.to_excel(writer, sheet_name="MASTER_MODELS", index=False)
            df_randomizer.to_excel(writer, sheet_name="ADVANCED_RANDOMIZER", index=False)
            
        print("Successfully updated V4 with ULTIMATE PROMPTS and RANDOMIZER.")
    else:
        print(f"File not found: {file_path}")
except PermissionError:
    print("PERMISSION_ERROR")
except Exception as e:
    print(f"Error: {e}")
